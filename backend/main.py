"""
SIF Safety Intelligence Platform - FastAPI Backend Server
Provides async batch report processing, SIF precursor detection,
prioritization, human review management, and dashboard analytics.
"""

import os
import uuid
import zipfile
import io
import asyncio
from datetime import datetime
from typing import List, Optional, Dict, Any

from fastapi import FastAPI, File, UploadFile, BackgroundTasks, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import JSONResponse
from pydantic import BaseModel

from backend.database import get_db_connection, init_db
from backend.analyzer import SafetyAnalyzer

app = FastAPI(
    title="SIF Safety Intelligence Platform API",
    description="Decision-support platform for Oil India Limited (OIL) safety reports & SIF precursor prioritization",
    version="1.0.0"
)

# Enable CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_DIR = os.getenv("UPLOAD_DIR", "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)

analyzer = SafetyAnalyzer()

# Pydantic Request Models
class ReviewUpdate(BaseModel):
    status: str # 'Pending Review', 'Confirmed', 'Not a SIF Precursor', 'Requires Further Review'
    comment: Optional[str] = ""
    reviewer: Optional[str] = "Safety Officer"

class AnalyzeRequest(BaseModel):
    report_ids: Optional[List[int]] = None # None means analyze all unanalyzed

@app.on_event("startup")
def on_startup():
    init_db()

# --- File Extraction Helper ---
def extract_text_from_file(filename: str, content: bytes) -> str:
    ext = filename.lower().split(".")[-1]
    if ext == "txt":
        try:
            return content.decode("utf-8")
        except UnicodeDecodeError:
            return content.decode("latin-1", errors="ignore")
    elif ext == "pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content))
            text = "\n".join([page.extract_text() or "" for page in reader.pages])
            return text.strip() or "[PDF contained no extractable text]"
        except Exception as e:
            return f"[Error extracting PDF: {str(e)}]"
    elif ext == "docx":
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            text = "\n".join([p.text for p in doc.paragraphs])
            return text.strip() or "[DOCX contained no text]"
        except Exception as e:
            return f"[Error extracting DOCX: {str(e)}]"
    else:
        return content.decode("utf-8", errors="ignore")

# --- Background Worker for Batch Processing ---
async def process_batch_job(job_id: str, report_ids: List[int]):
    conn = get_db_connection()
    cursor = conn.cursor()

    total = len(report_ids)
    completed = 0
    failed = 0

    for rep_id in report_ids:
        try:
            cursor.execute("SELECT id, filename, report_text FROM reports WHERE id = ?", (rep_id,))
            row = cursor.fetchone()
            if not row:
                continue

            fname = row["filename"]
            rtext = row["report_text"]

            # Update job progress
            cursor.execute(
                "UPDATE jobs SET current_file = ?, current_step = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                (fname, "Extracting report text & analyzing safety factors", job_id)
            )
            conn.commit()

            cursor.execute("UPDATE reports SET processing_status = 'processing' WHERE id = ?", (rep_id,))
            conn.commit()

            # Analyze report
            extraction = analyzer.extract_with_llm(rtext, fname)
            result = analyzer.assess_risk_and_prioritize(extraction, rtext)

            # Save Analysis
            cursor.execute("""
            INSERT INTO analysis (
                report_id, sif_detected, risk_score, risk_level, priority_rank,
                explanation, summary, worker_exposure, missing_controls, suggested_actions, analyzed_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(report_id) DO UPDATE SET
                sif_detected = excluded.sif_detected,
                risk_score = excluded.risk_score,
                risk_level = excluded.risk_level,
                priority_rank = excluded.priority_rank,
                explanation = excluded.explanation,
                summary = excluded.summary,
                worker_exposure = excluded.worker_exposure,
                missing_controls = excluded.missing_controls,
                suggested_actions = excluded.suggested_actions,
                analyzed_at = CURRENT_TIMESTAMP
            """, (
                rep_id, result.sif_detected, result.risk_score, result.risk_level, result.priority_rank,
                result.explanation, result.summary, 1 if result.worker_exposure else 0,
                ", ".join(result.missing_controls), "\n".join(result.suggested_actions)
            ))

            # Save Factors
            cursor.execute("DELETE FROM factors WHERE report_id = ?", (rep_id,))
            for f in result.factors:
                cursor.execute("""
                INSERT INTO factors (report_id, factor_type, factor_name, evidence)
                VALUES (?, ?, ?, ?)
                """, (rep_id, f.get("factor_type", "hazard"), f.get("factor_name", "Factor"), f.get("evidence", "")))

            # Initialize Review if not present
            cursor.execute("""
            INSERT OR IGNORE INTO reviews (report_id, status, comment, reviewer, reviewed_at)
            VALUES (?, 'Pending Review', '', 'Safety Officer', CURRENT_TIMESTAMP)
            """, (rep_id,))

            # Update report status and extracted metadata
            cursor.execute("""
            UPDATE reports SET
                report_type = ?,
                location = ?,
                activity = ?,
                equipment = ?,
                processing_status = 'completed'
            WHERE id = ?
            """, (result.report_type, result.location, result.activity, result.equipment, rep_id))

            completed += 1

        except Exception as e:
            failed += 1
            print(f"[Batch Processing Error] Failed to process report {rep_id}: {e}")
            cursor.execute("UPDATE reports SET processing_status = 'failed' WHERE id = ?", (rep_id,))

        cursor.execute(
            "UPDATE jobs SET completed_reports = ?, failed_reports = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            (completed, failed, job_id)
        )
        conn.commit()
        await asyncio.sleep(0.05) # Yield event loop

    cursor.execute(
        "UPDATE jobs SET status = 'completed', current_step = 'Batch analysis complete', updated_at = CURRENT_TIMESTAMP WHERE id = ?",
        (job_id,)
    )
    conn.commit()
    conn.close()

# --- REST Endpoints ---

@app.get("/api/health")
def get_health():
    return {
        "status": "healthy",
        "service": "SIF Safety Intelligence Platform",
        "llm_provider": analyzer.provider,
        "llm_model": analyzer.model_name,
        "timestamp": datetime.utcnow().isoformat()
    }

@app.post("/api/upload")
async def upload_reports(files: List[UploadFile] = File(...)):
    """Uploads multiple PDF, DOCX, TXT files, or a ZIP archive of reports."""
    conn = get_db_connection()
    cursor = conn.cursor()
    saved_ids = []

    for file in files:
        filename = file.filename or "unknown.txt"
        ext = filename.lower().split(".")[-1]
        content = await file.read()

        if ext == "zip":
            try:
                with zipfile.ZipFile(io.BytesIO(content)) as z:
                    for zname in z.namelist():
                        if zname.endswith("/") or zname.startswith("__MACOSX"):
                            continue
                        zext = zname.lower().split(".")[-1]
                        if zext in ["txt", "pdf", "docx"]:
                            zcontent = z.read(zname)
                            text = extract_text_from_file(zname, zcontent)
                            cursor.execute("""
                            INSERT INTO reports (filename, file_size, file_type, report_text, processing_status)
                            VALUES (?, ?, ?, ?, 'pending')
                            """, (os.path.basename(zname), len(zcontent), zext, text))
                            saved_ids.append(cursor.lastrowid)
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Error reading ZIP file: {str(e)}")
        elif ext in ["txt", "pdf", "docx"]:
            text = extract_text_from_file(filename, content)
            cursor.execute("""
            INSERT INTO reports (filename, file_size, file_type, report_text, processing_status)
            VALUES (?, ?, ?, ?, 'pending')
            """, (filename, len(content), ext, text))
            saved_ids.append(cursor.lastrowid)
        else:
            # Unsupported file type skipped or flagged
            continue

    conn.commit()
    conn.close()

    return {
        "message": f"Successfully uploaded and extracted {len(saved_ids)} reports.",
        "uploaded_count": len(saved_ids),
        "report_ids": saved_ids
    }

@app.post("/api/analyze")
async def start_analysis(req: AnalyzeRequest, background_tasks: BackgroundTasks):
    """Initiates asynchronous batch analysis for safety reports."""
    conn = get_db_connection()
    cursor = conn.cursor()

    if req.report_ids and len(req.report_ids) > 0:
        placeholders = ",".join("?" for _ in req.report_ids)
        cursor.execute(f"SELECT id FROM reports WHERE id IN ({placeholders})", req.report_ids)
    else:
        cursor.execute("SELECT id FROM reports WHERE processing_status != 'completed'")

    rows = cursor.fetchall()
    target_ids = [r["id"] for r in rows]

    if not target_ids:
        conn.close()
        return {"message": "No reports eligible for analysis", "job_id": None, "total": 0}

    job_id = str(uuid.uuid4())
    cursor.execute("""
    INSERT INTO jobs (id, total_reports, completed_reports, failed_reports, current_step, status)
    VALUES (?, ?, 0, 0, 'Initializing pipeline', 'running')
    """, (job_id, len(target_ids)))
    conn.commit()
    conn.close()

    background_tasks.add_task(process_batch_job, job_id, target_ids)

    return {
        "message": "Analysis started in background.",
        "job_id": job_id,
        "total": len(target_ids)
    }

@app.get("/api/jobs/{job_id}")
def get_job_status(job_id: str):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM jobs WHERE id = ?", (job_id,))
    job = cursor.fetchone()
    conn.close()

    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    return dict(job)

@app.get("/api/reports")
def get_reports(
    filter_by: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    sort_by: Optional[str] = Query("priority_rank"), # priority_rank, risk_score, date
    sort_order: Optional[str] = Query("asc")
):
    """Fetches reports with prioritization, risk levels, and human review status."""
    conn = get_db_connection()
    cursor = conn.cursor()

    query = """
    SELECT 
        r.id, r.filename, r.report_type, r.location, r.uploaded_at, r.processing_status,
        a.sif_detected, a.risk_score, a.risk_level, a.priority_rank, a.summary,
        rev.status as review_status, rev.reviewer, rev.reviewed_at
    FROM reports r
    LEFT JOIN analysis a ON r.id = a.report_id
    LEFT JOIN reviews rev ON r.id = rev.report_id
    WHERE 1=1
    """
    params = []

    # Filter conditions
    if filter_by:
        f = filter_by.lower()
        if f in ["critical", "high", "medium", "low"]:
            query += " AND LOWER(a.risk_level) = ?"
            params.append(f)
        elif f == "sif":
            query += " AND a.sif_detected = 'YES'"
        elif f in ["unsafe_act", "unsafe act"]:
            query += " AND (r.report_type = 'unsafe_act' OR r.report_type = 'Unsafe Act')"
        elif f in ["unsafe_condition", "unsafe condition"]:
            query += " AND (r.report_type = 'unsafe_condition' OR r.report_type = 'Unsafe Condition')"
        elif f in ["near_miss", "near miss"]:
            query += " AND (r.report_type = 'near_miss' OR r.report_type = 'Near Miss')"
        elif f == "pending":
            query += " AND (rev.status = 'Pending Review' OR rev.status IS NULL)"
        elif f == "reviewed":
            query += " AND rev.status IN ('Confirmed', 'Not a SIF Precursor', 'Requires Further Review')"

    if search:
        query += " AND (r.filename LIKE ? OR r.report_text LIKE ? OR a.summary LIKE ? OR r.location LIKE ?)"
        s = f"%{search}%"
        params.extend([s, s, s, s])

    # Sorting
    if sort_by == "risk_score":
        order_col = "COALESCE(a.risk_score, 0)"
        direction = "DESC" if sort_order.lower() == "desc" else "ASC"
    elif sort_by == "date":
        order_col = "r.uploaded_at"
        direction = "DESC" if sort_order.lower() == "desc" else "ASC"
    else:
        # Default priority ranking: Critical SIF first
        order_col = "COALESCE(a.priority_rank, 999) ASC, COALESCE(a.risk_score, 0) DESC"
        direction = ""

    query += f" ORDER BY {order_col} {direction}"

    cursor.execute(query, params)
    rows = cursor.fetchall()
    results = [dict(r) for r in rows]

    # Attach key factors snippet
    for r in results:
        cursor.execute("SELECT factor_name FROM factors WHERE report_id = ? LIMIT 3", (r["id"],))
        f_rows = cursor.fetchall()
        r["key_factors"] = [f["factor_name"] for f in f_rows]

    conn.close()
    return {"reports": results, "count": len(results)}

@app.get("/api/reports/{report_id}")
def get_report_detail(report_id: int):
    """Returns complete details, original text, safety factors, explanation, suggestions, and review logs."""
    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM reports WHERE id = ?", (report_id,))
    rep = cursor.fetchone()
    if not rep:
        conn.close()
        raise HTTPException(status_code=404, detail="Report not found")

    cursor.execute("SELECT * FROM analysis WHERE report_id = ?", (report_id,))
    ana = cursor.fetchone()

    cursor.execute("SELECT * FROM factors WHERE report_id = ?", (report_id,))
    factors = cursor.fetchall()

    cursor.execute("SELECT * FROM reviews WHERE report_id = ?", (report_id,))
    rev = cursor.fetchone()

    conn.close()

    return {
        "report": dict(rep),
        "analysis": dict(ana) if ana else None,
        "factors": [dict(f) for f in factors],
        "review": dict(rev) if rev else {"status": "Pending Review", "comment": "", "reviewer": ""}
    }

@app.patch("/api/reports/{report_id}/review")
def update_review(report_id: int, payload: ReviewUpdate):
    """Updates human safety officer review decision and timestamp."""
    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT id FROM reports WHERE id = ?", (report_id,))
    if not cursor.fetchone():
        conn.close()
        raise HTTPException(status_code=404, detail="Report not found")

    cursor.execute("""
    INSERT INTO reviews (report_id, status, comment, reviewer, reviewed_at)
    VALUES (?, ?, ?, ?, CURRENT_TIMESTAMP)
    ON CONFLICT(report_id) DO UPDATE SET
        status = excluded.status,
        comment = excluded.comment,
        reviewer = excluded.reviewer,
        reviewed_at = CURRENT_TIMESTAMP
    """, (report_id, payload.status, payload.comment, payload.reviewer))

    conn.commit()
    conn.close()

    return {"message": "Human review updated successfully.", "status": payload.status}

@app.get("/api/dashboard")
def get_dashboard_analytics():
    """Returns database-derived metrics, SIF distribution, hazard trends, and review statuses."""
    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT COUNT(*) as count FROM reports")
    total_reports = cursor.fetchone()["count"]

    cursor.execute("SELECT COUNT(*) as count FROM analysis WHERE sif_detected = 'YES'")
    sif_reports = cursor.fetchone()["count"]

    cursor.execute("SELECT COUNT(*) as count FROM analysis WHERE risk_level = 'CRITICAL'")
    critical = cursor.fetchone()["count"]

    cursor.execute("SELECT COUNT(*) as count FROM analysis WHERE risk_level = 'HIGH'")
    high = cursor.fetchone()["count"]

    cursor.execute("SELECT COUNT(*) as count FROM analysis WHERE risk_level = 'MEDIUM'")
    medium = cursor.fetchone()["count"]

    cursor.execute("SELECT COUNT(*) as count FROM analysis WHERE risk_level = 'LOW'")
    low = cursor.fetchone()["count"]

    # Reports by Type
    cursor.execute("""
    SELECT report_type, COUNT(*) as count
    FROM reports
    WHERE report_type IS NOT NULL AND report_type != ''
    GROUP BY report_type
    """)
    by_type = [dict(r) for r in cursor.fetchall()]

    # Top Hazards
    cursor.execute("""
    SELECT factor_name, COUNT(*) as count
    FROM factors
    WHERE factor_type = 'hazard'
    GROUP BY factor_name
    ORDER BY count DESC
    LIMIT 6
    """)
    top_hazards = [dict(r) for r in cursor.fetchall()]

    # Review status breakdown
    cursor.execute("""
    SELECT status, COUNT(*) as count
    FROM reviews
    GROUP BY status
    """)
    review_status = [dict(r) for r in cursor.fetchall()]

    conn.close()

    return {
        "total_reports": total_reports,
        "sif_reports": sif_reports,
        "critical": critical,
        "high": high,
        "medium": medium,
        "low": low,
        "reports_by_type": by_type,
        "top_hazards": top_hazards,
        "review_status": review_status
    }

@app.post("/api/reset")
@app.post("/api/clear")
@app.post("/api/database/reset")
def reset_database_api():
    """Reset database tables and verify zero state."""
    from backend.database import reset_db
    counts = reset_db()
    return {
        "message": "Database successfully reset to zero state.",
        "database_path": DB_PATH,
        "row_counts": counts
    }

